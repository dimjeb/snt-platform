"""
Разбор банковской выписки в виде Excel-таблицы.

Формат «1С:Клиент-Банк» (`statement.py`) остаётся основным — он
стабилен и одинаков у всех банков. Но ВТБ Бизнес по умолчанию отдаёт
кнопкой «Выписка» именно .xlsx, и председатель приносит то, что выгрузил
банк, а не то, что удобно нам. Поэтому файл разбираем и такой.

Опираемся не на номера колонок, а на подписи в шапке: банк без
предупреждения добавляет налоговые поля справа и меняет порядок.
Ищем строку, где одновременно есть «Дата» и «Назначение», от неё и
пляшем.

Устройство выгрузки ВТБ:

    строка 1  ВЫПИСКА
    строка 2  Номер счета: 40703810100810020382 | Валюта: ... | Владелец счёта: ...
    строка 3  Начальная дата: 31.08.2026 | Конечная дата: 29.09.2026
    строка 4  Входящий остаток RUB: 0 | Исходящий остаток RUB: 0
    строка 7  Дата | Номер | Вид операции | Контрагент | ИНН контрагента |
              БИК банка контрагента | Счет контрагента | Дебет, RUR |
              Кредит, RUR | Назначение | ...налоговые поля...
    строка 8+ операции, между днями — строки «ИТОГО ЗА ДЕНЬ:», в конце «ИТОГО:»

Приход — это строки, где заполнена колонка «Кредит». Расход нам не
нужен: разносить по начислениям нечего.
"""
import re
from datetime import date, datetime
from decimal import Decimal

from .statement import ParsedStatement, StatementDocument, StatementError, _parse_amount

# Первые байты zip-архива. .xlsx — это zip, и по расширению судить
# нельзя: файл приходит из браузера с тем именем, какое дал банк.
XLSX_MAGIC = b"PK\x03\x04"
# Старый бинарный .xls (OLE2). Прочитать его openpyxl не может, и надо
# сказать об этом человеку понятно, а не «файл повреждён».
XLS_MAGIC = b"\xd0\xcf\x11\xe0"

MAX_HEADER_SCAN = 30       # в скольких верхних строках искать шапку
TOTAL_MARKERS = ("итого", "всего", "оборот")


def looks_like_xlsx(raw: bytes) -> bool:
    return raw[:4] == XLSX_MAGIC


def looks_like_old_xls(raw: bytes) -> bool:
    return raw[:4] == XLS_MAGIC


def _text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return value.strftime("%d.%m.%Y")
    return str(value).strip()


def _norm(value) -> str:
    """Подпись колонки без пробелов, регистра и хвостовых знаков."""
    return re.sub(r"[\s ]+", " ", _text(value)).strip(" :").lower()


def _cell_date(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = _text(value)
    for fmt in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y", "%d.%m.%y"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def _cell_amount(value) -> Decimal:
    if isinstance(value, (int, float, Decimal)):
        return Decimal(str(value))
    return _parse_amount(_text(value))


def _clean_payer(value: str) -> str:
    """
    Достать ФИО плательщика из строки контрагента.

    СБП-переводы банк пишет одной строкой с разделителем «//»:

        ПАО СБЕРБАНК//ИВАНОВ ИВАН ИВАНОВИЧ//3600467115214//381505107980//

    Сопоставление по ФИО (`matching._name_key`) берёт первое слово как
    фамилию, так что без разбора каждый такой платёж «платил бы ПАО
    Сбербанк». Берём часть, похожую на ФИО: два и более слова из букв.
    """
    text = _text(value)
    if "//" not in text:
        return text
    parts = [p.strip() for p in text.split("//") if p.strip()]
    for part in parts[1:]:
        words = [w for w in re.split(r"\s+", part) if re.fullmatch(r"[^\W\d_]+", w)]
        if len(words) >= 2:
            return part
    return parts[0] if parts else text


def _find_header(rows):
    """Строка шапки таблицы: в ней есть и «Дата», и «Назначение»."""
    for index, row in enumerate(rows[:MAX_HEADER_SCAN]):
        titles = {_norm(cell) for cell in row}
        if "дата" in titles and any(t.startswith("назначение") for t in titles):
            return index
    return None


def _columns(header_row):
    """Подпись колонки → её индекс."""
    found = {}
    for index, cell in enumerate(header_row):
        title = _norm(cell)
        if not title:
            continue
        if title == "дата":
            found.setdefault("date", index)
        elif title == "номер":
            found.setdefault("number", index)
        elif title.startswith("контрагент"):
            found.setdefault("payer", index)
        elif title.startswith("инн"):
            found.setdefault("inn", index)
        elif title.startswith("счет контрагента") or title.startswith("счёт контрагента"):
            found.setdefault("payer_account", index)
        elif title.startswith("дебет"):
            found.setdefault("debit", index)
        elif title.startswith("кредит"):
            found.setdefault("credit", index)
        elif title.startswith("назначение"):
            found.setdefault("purpose", index)
    return found


def _header_value(rows, header_index, *labels):
    """
    Значение из шапки файла: ищем подпись и берём соседнюю клетку справа.

    Подписи в выгрузке лежат не в фиксированных клетках: «Номер счета:»
    стоит в A2, а «Владелец счёта:» — в F2.
    """
    for row in rows[:header_index]:
        for index, cell in enumerate(row):
            title = _norm(cell)
            if any(title.startswith(label) for label in labels):
                for value in row[index + 1:index + 3]:
                    if _text(value):
                        return _text(value)
    return ""


def parse_xlsx_statement(raw: bytes, *, our_account: str = "") -> ParsedStatement:
    """
    Разобрать выписку-таблицу.

    our_account принимается ради единой сигнатуры с parse_1c_statement,
    но здесь не фильтрует: в таблице нет колонки «получатель» — выписка
    целиком по одному счёту, и он в шапке. Что счёт наш, проверяет
    import_statement, сверяя result.account с реквизитами организации.
    """
    import io

    import openpyxl

    try:
        book = openpyxl.load_workbook(io.BytesIO(raw), data_only=True, read_only=True)
    except Exception as exc:                                  # noqa: BLE001
        raise StatementError(
            "Не удалось открыть файл как таблицу Excel. "
            f"Банк-клиент выгрузил что-то другое ({exc.__class__.__name__})."
        ) from exc

    sheet = book.worksheets[0]
    rows = [list(row) for row in sheet.iter_rows(values_only=True)]
    book.close()

    header_index = _find_header(rows)
    if header_index is None:
        raise StatementError(
            "В таблице нет шапки с колонками «Дата» и «Назначение» — "
            "это не выписка. Выгрузите из банк-клиента выписку по счёту "
            "за период или файл обмена с 1С."
        )

    columns = _columns(rows[header_index])
    required = {"date": "Дата", "credit": "Кредит", "purpose": "Назначение"}
    missing = [title for key, title in required.items() if key not in columns]
    if missing:
        raise StatementError(
            "В выписке не хватает колонок: " + ", ".join(f"«{t}»" for t in missing) + "."
        )

    result = ParsedStatement()
    result.account = re.sub(r"\D", "", _header_value(rows, header_index, "номер счета", "номер счёта", "счет", "счёт"))
    if not result.account and re.fullmatch(r"\d{20}", sheet.title.strip()):
        # ВТБ называет лист номером счёта — запасной источник.
        result.account = sheet.title.strip()
    result.date_from = _cell_date(_header_value(rows, header_index, "начальная дата", "дата начала", "период с"))
    result.date_to = _cell_date(_header_value(rows, header_index, "конечная дата", "дата окончания", "период по"))
    result.sender = _header_value(rows, header_index, "владелец сч")

    def cell(row, key):
        index = columns.get(key)
        if index is None or index >= len(row):
            return None
        return row[index]

    for row in rows[header_index + 1:]:
        first = _norm(cell(row, "date"))
        if any(first.startswith(marker) for marker in TOTAL_MARKERS):
            # «ИТОГО ЗА ДЕНЬ» и «ИТОГО» — подбивка, а не платёж.
            continue
        doc_date = _cell_date(cell(row, "date"))
        if doc_date is None:
            continue
        amount = _cell_amount(cell(row, "credit"))
        if amount <= 0:
            # Пустой день или списание — разносить нечего.
            continue
        result.documents.append(StatementDocument(
            number=_text(cell(row, "number")),
            doc_date=doc_date,
            amount=amount,
            payer_name=_clean_payer(_text(cell(row, "payer"))),
            payer_account=re.sub(r"\D", "", _text(cell(row, "payer_account"))),
            payer_inn=re.sub(r"\D", "", _text(cell(row, "inn"))),
            payee_account=result.account,
            purpose=_text(cell(row, "purpose")),
            raw={"contragent": _text(cell(row, "payer"))},
        ))

    return result
