"""
Чтение таблиц Excel для загрузки данных с сайта.

Колонки узнаются по заголовкам, а не по порядку: люди переставляют
столбцы, добавляют свои («Примечание», «Улица») — и это не должно ломать
загрузку. Заголовок ищется в первых строках листа: над таблицей часто
стоит название документа или дата.

Ни содержимое строк, ни имя файла в журнал не пишутся: в реестре —
персональные данные.
"""
import datetime as dt
import io
from decimal import Decimal, InvalidOperation

from rest_framework import serializers

MAX_UPLOAD_BYTES = 5 * 1024 * 1024
HEADER_SCAN_ROWS = 10
XLSX_MAGIC = b"PK\x03\x04"


class ImportFileError(Exception):
    """Файл не читается как ожидаемая таблица — сообщение для человека."""


def read_upload(request):
    """Достать загруженный .xlsx из запроса или объяснить, что не так."""
    f = request.FILES.get("file")
    if f is None:
        raise serializers.ValidationError({"file": "Выберите файл Excel (.xlsx)."})
    if f.size > MAX_UPLOAD_BYTES:
        raise serializers.ValidationError(
            {"file": "Файл больше 5 МБ. Для реестра и показаний этого с запасом — "
                     "проверьте, тот ли файл."})
    head = f.read(4)
    f.seek(0)
    if head != XLSX_MAGIC:
        raise serializers.ValidationError(
            {"file": "Нужен файл .xlsx. Старый .xls или .csv откройте в Excel "
                     "и сохраните как «Книга Excel (.xlsx)»."})
    return f


def is_dry_run(request):
    value = request.data.get("dry_run", "true")
    return str(value).lower() in ("1", "true", "yes", "on")


def _norm(text):
    return " ".join(str(text or "").replace("\xa0", " ").lower()
                    .replace("ё", "е").split())


def read_table(fileobj, columns, required):
    """
    Прочитать первый лист.

    columns — упорядоченный {ключ: (подстроки заголовка, ...)}. Порядок
    важен: «доп. телефон» содержит «телефон», поэтому более узкий ключ
    должен стоять раньше. Каждой колонке листа достаётся первый
    подошедший ключ, каждому ключу — первая подошедшая колонка.

    Возвращает список (номер строки в Excel, {ключ: значение}) — только
    строки, где есть хоть что-то.
    """
    import openpyxl

    try:
        wb = openpyxl.load_workbook(io.BytesIO(fileobj.read()),
                                    read_only=True, data_only=True)
    except Exception:
        raise ImportFileError("Файл не открывается как таблица Excel (.xlsx).")
    try:
        ws = wb.worksheets[0]
        rows = list(ws.iter_rows(values_only=True))
    finally:
        wb.close()

    header_idx, mapping = None, {}
    for idx, row in enumerate(rows[:HEADER_SCAN_ROWS]):
        found = {}
        for col, cell in enumerate(row):
            text = _norm(cell)
            if not text:
                continue
            for key, needles in columns.items():
                if key in found:
                    continue
                if any(n in text for n in needles):
                    found[key] = col
                    break
        if required <= set(found):
            header_idx, mapping = idx, found
            break
    if header_idx is None:
        names = ", ".join(f"«{columns[k][0]}»" for k in columns if k in required)
        raise ImportFileError(
            f"Не нашёл строку заголовков. В первых {HEADER_SCAN_ROWS} строках "
            f"листа должны быть колонки: {names}. Возьмите шаблон — в нём всё "
            f"подписано.")

    result = []
    for offset, row in enumerate(rows[header_idx + 1:], start=header_idx + 2):
        values = {key: (row[col] if col < len(row) else None)
                  for key, col in mapping.items()}
        if any(v is not None and str(v).strip() != "" for v in values.values()):
            result.append((offset, values))
    return result


def cell_text(value):
    """Текст ячейки; число 12.0 из Excel становится «12», а не «12.0»."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).replace("\xa0", " ").strip()


def cell_decimal(value):
    """Число из ячейки или None; «1 250,50» тоже число."""
    text = cell_text(value).replace(" ", "").replace(",", ".")
    if not text:
        return None
    try:
        return Decimal(text)
    except InvalidOperation:
        raise ValueError(text)


def cell_date(value):
    """Дата из ячейки: настоящая дата Excel или «дд.мм.гггг»."""
    if value is None or cell_text(value) == "":
        return None
    if isinstance(value, dt.datetime):
        return value.date()
    if isinstance(value, dt.date):
        return value
    text = cell_text(value)
    for fmt in ("%d.%m.%Y", "%d.%m.%y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return dt.datetime.strptime(text, fmt).date()
        except ValueError:
            pass
    raise ValueError(text)


def cell_flag(value):
    """«да», «1», «+», «x» — да; всё остальное, включая пусто, — нет."""
    return _norm(cell_text(value)) in ("да", "1", "+", "x", "х", "true", "yes", "v")


def template_response(filename, header, examples, notes):
    """Шаблон .xlsx: лист с заголовком и примерами + лист с пояснениями."""
    import openpyxl
    from django.http import HttpResponse
    from openpyxl.styles import Font
    from urllib.parse import quote

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Данные"
    ws.append(header)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in examples:
        ws.append(row)
    for idx, title in enumerate(header, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(idx)].width = \
            max(14, len(str(title)) + 4)
    help_ws = wb.create_sheet("Пояснения")
    for line in notes:
        help_ws.append([line])
    help_ws.column_dimensions["A"].width = 110

    buf = io.BytesIO()
    wb.save(buf)
    resp = HttpResponse(
        buf.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    resp["Content-Disposition"] = (
        f"attachment; filename=template.xlsx; filename*=UTF-8''{quote(filename)}")
    return resp


class _Rollback(Exception):
    """Откатить транзакцию проверочного прогона."""


def run_import(fn, dry_run):
    """
    Выполнить импорт в транзакции; в режиме проверки — откатить.

    Проверка идёт тем же кодом, что и запись: отчёт «что будет» не может
    разойтись с тем, что получится на самом деле.
    """
    from django.db import transaction

    result = {}
    try:
        with transaction.atomic():
            result = fn()
            if dry_run:
                raise _Rollback()
    except _Rollback:
        pass
    return result


def import_response(request, columns, required, do_import):
    """
    Общая часть ручек «Загрузить из Excel»: файл → таблица → импорт
    (или проверка без записи) → отчёт.

    do_import(rows) возвращает {"rows", "stats", "issues"}.
    """
    from rest_framework.response import Response

    f = read_upload(request)
    dry_run = is_dry_run(request)
    try:
        rows = read_table(f, columns, required)
    except ImportFileError as exc:
        raise serializers.ValidationError({"file": str(exc)})
    if not rows:
        raise serializers.ValidationError(
            {"file": "Под заголовками нет ни одной заполненной строки."})
    result = run_import(lambda: do_import(rows), dry_run)
    return Response({"dry_run": dry_run, **result})
