"""
Документы Excel: смета, финансово-экономическое обоснование, отчёт об
исполнении сметы, ведомость для ревизионной комиссии, выгрузка для
бухгалтера.

Документы для собрания оформлены так, чтобы их можно было распечатать
и подписать как есть: шапка с наименованием товарищества, строки
«Утверждено решением общего собрания» и подписи.
"""
import io
from collections import defaultdict
from decimal import Decimal

import openpyxl
from django.db.models import Sum
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ZERO = Decimal("0")
MONEY = '#,##0.00'
THIN = Side(style="thin", color="999999")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEAD_FILL = PatternFill(fill_type="solid", fgColor="E2EFDA")
CATEGORY_LABELS = {"membership": "Членский взнос", "target": "Целевой взнос",
                   "electricity": "Электроэнергия", "penalty": "Пени",
                   "other": "Прочее"}


def _org_name(org):
    return org.full_name or org.name


def _approval_line(budget):
    if budget.is_approved and budget.approved_at:
        proto = f", протокол № {budget.protocol_number}" if budget.protocol_number else ""
        return (f"УТВЕРЖДЕНА решением общего собрания членов товарищества "
                f"от {budget.approved_at:%d.%m.%Y}{proto}")
    return "УТВЕРЖДЕНА решением общего собрания членов товарищества от «___» ________ 20__ г., протокол № ____"


class Sheet:
    """Тонкая обёртка: строки подряд, заголовки таблиц, деньги, ширины."""

    def __init__(self, ws, widths):
        self.ws = ws
        self.row = 1
        for idx, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(idx)].width = w
        self.cols = len(widths)

    def line(self, text, bold=False, size=11, align="left"):
        c = self.ws.cell(row=self.row, column=1, value=text)
        c.font = Font(bold=bold, size=size)
        c.alignment = Alignment(horizontal=align, wrap_text=True, vertical="top")
        self.ws.merge_cells(start_row=self.row, start_column=1,
                            end_row=self.row, end_column=self.cols)
        self.row += 1

    def blank(self, n=1):
        self.row += n

    def header(self, titles):
        for idx, t in enumerate(titles, start=1):
            c = self.ws.cell(row=self.row, column=idx, value=t)
            c.font = Font(bold=True)
            c.fill = HEAD_FILL
            c.border = BOX
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        self.row += 1

    def values(self, values, bold=False, money=()):
        for idx, v in enumerate(values, start=1):
            c = self.ws.cell(row=self.row, column=idx, value=v)
            c.border = BOX
            c.alignment = Alignment(vertical="top", wrap_text=True)
            if bold:
                c.font = Font(bold=True)
            if idx in money and v is not None and v != "":
                c.number_format = MONEY
        self.row += 1


def _save(wb):
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _signatures(sh, org):
    sh.blank()
    sh.line("Председатель правления  _______________ / _______________ /")
    sh.blank()
    sh.line("Казначей  _______________ / _______________ /")


# ── смета ─────────────────────────────────────────────────────────────────────

def smeta_xlsx(budget, calc):
    from .models import Budget, BudgetItem

    org = budget.organization
    wb = openpyxl.Workbook()
    sh = Sheet(wb.active, [6, 60, 18])
    wb.active.title = "Смета"
    sh.line(_approval_line(budget), size=9, align="right")
    sh.blank()
    sh.line(_org_name(org), bold=True, align="center")
    sh.line(f"ПРИХОДНО-РАСХОДНАЯ СМЕТА на {budget.year} год", bold=True, size=13,
            align="center")
    sh.blank()

    sh.line("1. Доходы", bold=True)
    sh.header(["№", "Источник", "Сумма, ₽"])
    if budget.basis == Budget.BASIS_PER_SOTKA:
        how = (f"{calc['area']} сот. × {calc['rate']} ₽" if calc["rate"] is not None
               else "площадь участков не заполнена")
    else:
        how = (f"{calc['plots']} уч. × {calc['rate']} ₽" if calc["rate"] is not None
               else "участков с собственником нет")
    sh.values([1, f"Членские взносы ({how})", calc["collected"] or ZERO], money=(3,))
    n = 1
    for t in calc["targets"]:
        n += 1
        sh.values([n, f"Целевой взнос: {t['name']}", t["amount"]], money=(3,))
    income = (calc["collected"] or ZERO) + calc["target_total"]
    sh.values(["", "Итого доходов", income], bold=True, money=(3,))
    sh.blank()

    sh.line("2. Расходы", bold=True)
    sh.header(["№", "Статья расходов", "Сумма, ₽"])
    n = 0
    for section, title in BudgetItem.SECTION_CHOICES:
        items = [i for i in budget.items.all() if i.section == section]
        if not items:
            continue
        sh.values(["", title, ""], bold=True)
        for i in items:
            n += 1
            sh.values([n, i.name, i.amount], money=(3,))
    total = calc["membership_total"] + calc["target_total"]
    sh.values(["", "Итого расходов", total], bold=True, money=(3,))
    sh.blank()
    if budget.basis == Budget.BASIS_PER_SOTKA:
        sh.line(f"Размер членского взноса: {calc['rate']} ₽ за сотку площади участка в год.")
    else:
        sh.line(f"Размер членского взноса: {calc['rate']} ₽ с участка в год.")
    if calc["rounding_surplus"]:
        sh.line(f"Превышение доходов над расходами за счёт округления ставки до "
                f"копейки: {calc['rounding_surplus']:.2f} ₽.", size=9)
    _signatures(sh, org)
    return _save(wb)


# ── финансово-экономическое обоснование ───────────────────────────────────────

def feo_xlsx(budget, calc):
    from .models import Budget, BudgetItem

    org = budget.organization
    wb = openpyxl.Workbook()
    sh = Sheet(wb.active, [5, 38, 26, 16, 45])
    wb.active.title = "ФЭО"
    sh.line(_approval_line(budget), size=9, align="right")
    sh.blank()
    sh.line(_org_name(org), bold=True, align="center")
    sh.line("ФИНАНСОВО-ЭКОНОМИЧЕСКОЕ ОБОСНОВАНИЕ", bold=True, size=13, align="center")
    sh.line(f"размера членских и целевых взносов на {budget.year} год", bold=True,
            align="center")
    sh.line("(ст. 14 Федерального закона от 29.07.2017 № 217-ФЗ)", size=9, align="center")
    sh.blank()

    for section, title in BudgetItem.SECTION_CHOICES:
        items = [i for i in budget.items.all() if i.section == section]
        if not items:
            continue
        sh.line(title, bold=True)
        sh.header(["№", "Статья", "Расчёт", "Сумма, ₽", "Обоснование"])
        for n, i in enumerate(items, start=1):
            if i.quantity is not None and i.unit_price is not None:
                q = i.quantity.normalize()
                calc_text = f"{q:f} {i.unit or ''} × {i.unit_price} ₽".replace("  ", " ")
            else:
                calc_text = "—"
            sh.values([n, i.name, calc_text, i.amount, i.justification], money=(4,))
        sh.values(["", "Итого", "", sum((i.amount for i in items), ZERO), ""],
                  bold=True, money=(4,))
        sh.blank()

    sh.line("Расчёт размера членского взноса", bold=True)
    total = calc["membership_total"]
    if budget.basis == Budget.BASIS_PER_SOTKA:
        sh.line(f"Расходы за счёт членских взносов {total:.2f} ₽ ÷ общая площадь "
                f"участков {calc['area']} сот. = {calc['rate']} ₽ за сотку в год "
                "(с округлением вверх до копейки).")
        sh.line(f"Взнос с участка = ставка × площадь участка. Учтено участков: "
                f"{calc['plots_with_area']}.")
    else:
        sh.line(f"Расходы за счёт членских взносов {total:.2f} ₽ ÷ количество "
                f"участков {calc['plots']} = {calc['rate']} ₽ с участка в год "
                "(с округлением вверх до копейки).")
    if calc["targets"]:
        sh.blank()
        sh.line("Расчёт целевых взносов", bold=True)
        for t in calc["targets"]:
            sh.line(f"«{t['name']}»: {t['amount']:.2f} ₽ ÷ {calc['plots']} уч. = "
                    f"{t['per_plot']} ₽ с участка.")
    for w in calc["warnings"]:
        sh.line("Внимание: " + w, size=9)
    _signatures(sh, org)
    return _save(wb)


# ── исполнение ───────────────────────────────────────────────────────────────

def execution_xlsx(budget, ex):
    from .models import BudgetItem

    org = budget.organization
    wb = openpyxl.Workbook()
    sh = Sheet(wb.active, [5, 46, 16, 16, 16, 10])
    wb.active.title = "Исполнение"
    sh.line(_org_name(org), bold=True, align="center")
    sh.line(f"ОТЧЁТ ОБ ИСПОЛНЕНИИ ПРИХОДНО-РАСХОДНОЙ СМЕТЫ за {budget.year} год",
            bold=True, size=13, align="center")
    sh.blank()
    sh.line("1. Доходы", bold=True)
    sh.header(["№", "Источник", "Начислено, ₽", "Поступило, ₽", "Долг, ₽", ""])
    for n, r in enumerate(ex["income"], start=1):
        sh.values([n, r["name"], r["charged"], r["paid"], r["debt"], ""], money=(3, 4, 5))
    sh.blank()
    sh.line("2. Расходы", bold=True)
    sh.header(["№", "Статья", "План, ₽", "Факт, ₽", "Отклонение, ₽", "%"])
    labels = dict(BudgetItem.SECTION_CHOICES)
    n = 0
    for section in labels:
        rows = [r for r in ex["rows"] if r["section"] == section]
        if not rows:
            continue
        sh.values(["", labels[section], "", "", "", ""], bold=True)
        for r in rows:
            n += 1
            sh.values([n, r["name"], r["plan"], r["fact"], r["diff"],
                       float(r["percent"]) if r["percent"] is not None else ""],
                      money=(3, 4, 5))
    if ex["outside_budget"]:
        sh.values(["", "Расходы вне статей сметы", "", ex["outside_budget"], "", ""],
                  money=(4,))
    sh.values(["", "Итого", ex["plan_total"], ex["fact_total"] + ex["outside_budget"],
               ex["plan_total"] - ex["fact_total"] - ex["outside_budget"], ""],
              bold=True, money=(3, 4, 5))
    _signatures(sh, org)
    return _save(wb)


# ── ведомость для ревизионной комиссии ───────────────────────────────────────

def _owner_names(org):
    from members.models import PlotOwnership

    names = defaultdict(list)
    for o in (PlotOwnership.objects.filter(organization=org, date_to__isnull=True)
              .select_related("member", "plot")):
        names[o.plot.number].append(o.member.full_name)
    return {k: ", ".join(v) for k, v in names.items()}


def revision_xlsx(org, year):
    """Расчёты с каждым участком за год — по видам взносов."""
    from billing.models import Charge, Payment
    from billing.services import _plot_sort_key

    charges = Charge.objects.filter(organization=org, period__year=year)
    charged = defaultdict(lambda: defaultdict(lambda: ZERO))
    for row in (charges.values("plot__number", "charge_type__category")
                .annotate(s=Sum("amount"))):
        charged[row["plot__number"]][row["charge_type__category"]] += row["s"] or ZERO
    paid = defaultdict(lambda: defaultdict(lambda: ZERO))
    for row in (Payment.objects.filter(charge__in=charges, is_cancelled=False)
                .values("charge__plot__number", "charge__charge_type__category")
                .annotate(s=Sum("amount"))):
        paid[row["charge__plot__number"]][row["charge__charge_type__category"]] += row["s"] or ZERO

    cats = ["membership", "target", "electricity", "penalty", "other"]
    used = [c for c in cats if any(charged[p][c] for p in charged)]
    owners = _owner_names(org)

    wb = openpyxl.Workbook()
    widths = [10, 34] + [15, 15] * len(used) + [15, 15, 15]
    sh = Sheet(wb.active, widths)
    wb.active.title = "Ведомость"
    sh.line(_org_name(org), bold=True, align="center")
    sh.line(f"ВЕДОМОСТЬ РАСЧЁТОВ С ЧЛЕНАМИ ТОВАРИЩЕСТВА за {year} год", bold=True,
            size=13, align="center")
    sh.line("для ревизионной комиссии", align="center")
    sh.blank()
    titles = ["Участок", "Собственник"]
    for c in used:
        titles += [f"{CATEGORY_LABELS[c]}: начислено", f"{CATEGORY_LABELS[c]}: оплачено"]
    titles += ["Всего начислено", "Всего оплачено", "Долг"]
    sh.header(titles)
    totals = defaultdict(lambda: ZERO)
    for number in sorted(charged, key=_plot_sort_key):
        row = [number, owners.get(number, "")]
        ch_sum = pd_sum = ZERO
        for c in used:
            row += [charged[number][c], paid[number][c]]
            ch_sum += charged[number][c]
            pd_sum += paid[number][c]
            totals[c + "_c"] += charged[number][c]
            totals[c + "_p"] += paid[number][c]
        row += [ch_sum, pd_sum, ch_sum - pd_sum]
        totals["c"] += ch_sum
        totals["p"] += pd_sum
        sh.values(row, money=tuple(range(3, len(row) + 1)))
    total_row = ["Итого", ""]
    for c in used:
        total_row += [totals[c + "_c"], totals[c + "_p"]]
    total_row += [totals["c"], totals["p"], totals["c"] - totals["p"]]
    sh.values(total_row, bold=True, money=tuple(range(3, len(total_row) + 1)))
    wb.active.freeze_panes = "C6"
    return _save(wb)


# ── выгрузка для бухгалтера ──────────────────────────────────────────────────

def accounting_xlsx(org, date_from, date_to):
    """
    Всё, что бухгалтеру нужно перенести в учёт за период, по листам:
    поступления (реальные деньги), начисления, расходы, сальдо по участкам.

    «Поступления» — только живые деньги: строки выписки и приём кассы. Зачёт
    аванса и перенос между начислениями — внутренние движения, денег они
    не приносят, и в приход их ставить нельзя, иначе доход задвоится.
    """
    from billing.models import BankTransaction, Charge, Payment, PlotCredit, Receipt
    from billing.services import _plot_sort_key

    from .models import Expense

    owners = _owner_names(org)
    method_labels = dict(Payment.METHOD_CHOICES)
    wb = openpyxl.Workbook()

    # 1. Поступления
    ws = wb.active
    ws.title = "Поступления"
    sh = Sheet(ws, [12, 14, 10, 30, 16, 18, 50, 50])
    sh.line(f"{_org_name(org)} — поступления денежных средств "
            f"с {date_from:%d.%m.%Y} по {date_to:%d.%m.%Y}", bold=True)
    sh.header(["Дата", "Документ", "Участок", "Плательщик / собственник", "Сумма, ₽",
               "Способ", "Назначение / примечание", "Куда разнесено"])
    rows = []
    for t in (BankTransaction.objects.filter(organization=org,
                                              status=BankTransaction.STATUS_APPLIED,
                                              date__range=(date_from, date_to))
              .select_related("plot")):
        rows.append((t.date, t.doc_number, t.plot.number if t.plot else "",
                     t.payer_name, t.amount, "Банк (выписка)", t.purpose, t.note))
    for r in (Receipt.objects.filter(organization=org, date__range=(date_from, date_to))
              .select_related("plot")):
        rows.append((r.date, f"касса №{r.pk}", r.plot.number,
                     owners.get(r.plot.number, ""), r.amount,
                     method_labels.get(r.method, r.method), r.note, r.allocation))
    # Оплаты, внесённые до журнала поступлений: не из выписки, не перенос,
    # не зачёт аванса и не привязанные к записи журнала.
    spends = PlotCredit.objects.filter(organization=org, payment__isnull=False) \
        .values_list("payment", flat=True)
    legacy = (Payment.objects.filter(organization=org, is_cancelled=False,
                                     date__range=(date_from, date_to),
                                     bank_transaction__isnull=True)
              .exclude(method=Payment.METHOD_TRANSFER)
              .exclude(external_ref__startswith="receipt-")
              .exclude(external_ref__startswith="transfer-")
              .exclude(pk__in=spends)
              .exclude(notes__startswith="Выписка")
              .select_related("charge__plot", "charge__charge_type", "charge__period"))
    for p in legacy:
        number = p.charge.plot.number
        rows.append((p.date, "оплата", number, owners.get(number, ""), p.amount,
                     method_labels.get(p.method, p.method), p.notes,
                     f"{p.charge.charge_type.name} {p.charge.period}"))
    rows.sort(key=lambda r: (r[0], _plot_sort_key(r[2] or "")))
    for r in rows:
        sh.values([r[0], r[1], r[2], r[3], r[4], r[5], r[6], r[7]], money=(5,))
    sh.values(["Итого", "", "", "", sum((r[4] for r in rows), ZERO), "", "", ""],
              bold=True, money=(5,))

    # 2. Начисления
    sh = Sheet(wb.create_sheet("Начисления"), [12, 12, 10, 30, 22, 16, 14, 40])
    sh.line(f"Начисления, созданные с {date_from:%d.%m.%Y} по {date_to:%d.%m.%Y}",
            bold=True)
    sh.header(["Дата", "Период", "Участок", "Собственник", "Вид", "Сумма, ₽",
               "Оплатить до", "Описание"])
    total = ZERO
    for c in (Charge.objects.filter(organization=org,
                                    created_at__date__range=(date_from, date_to))
              .select_related("plot", "period", "charge_type")
              .order_by("created_at", "pk")):
        sh.values([c.created_at.date(), str(c.period), c.plot.number,
                   owners.get(c.plot.number, ""), c.charge_type.name, c.amount,
                   c.due_date or "", c.description], money=(6,))
        total += c.amount
    sh.values(["Итого", "", "", "", "", total, "", ""], bold=True, money=(6,))

    # 3. Расходы
    sh = Sheet(wb.create_sheet("Расходы"), [12, 16, 36, 30, 30, 40])
    sh.line(f"Расходы с {date_from:%d.%m.%Y} по {date_to:%d.%m.%Y}", bold=True)
    sh.header(["Дата", "Сумма, ₽", "Статья сметы", "Кому", "Документ", "Описание"])
    total = ZERO
    for e in (Expense.objects.filter(organization=org, date__range=(date_from, date_to))
              .select_related("item").order_by("date", "pk")):
        sh.values([e.date, e.amount, e.item.name if e.item else "вне сметы",
                   e.counterparty, e.document, e.description], money=(2,))
        total += e.amount
    sh.values(["Итого", total, "", "", "", ""], bold=True, money=(2,))

    # 4. Сальдо по участкам
    sh = Sheet(wb.create_sheet("Сальдо по участкам"), [10, 34, 16, 16, 16, 16, 16])
    sh.line(f"Расчёты с участками: долг на {date_from:%d.%m.%Y} и на "
            f"{date_to:%d.%m.%Y}", bold=True)
    sh.header(["Участок", "Собственник", "Долг на начало, ₽", "Начислено, ₽",
               "Оплачено, ₽", "Долг на конец, ₽", "Аванс на конец, ₽"])

    def by_plot(qs, field, key):
        return {r[key]: r["s"] or ZERO for r in qs.values(key).annotate(s=Sum(field))}

    ch = Charge.objects.filter(organization=org)
    pay = Payment.objects.filter(organization=org, is_cancelled=False)
    ch_before = by_plot(ch.filter(created_at__date__lt=date_from), "amount", "plot__number")
    ch_in = by_plot(ch.filter(created_at__date__range=(date_from, date_to)), "amount",
                    "plot__number")
    pay_before = by_plot(pay.filter(date__lt=date_from), "amount", "charge__plot__number")
    pay_in = by_plot(pay.filter(date__range=(date_from, date_to)), "amount",
                     "charge__plot__number")
    adv = by_plot(PlotCredit.objects.filter(organization=org, date__lte=date_to),
                  "amount", "plot__number")
    numbers = set(ch_before) | set(ch_in) | set(pay_before) | set(pay_in) | set(adv)
    tot = defaultdict(lambda: ZERO)
    for number in sorted(numbers, key=_plot_sort_key):
        start = ch_before.get(number, ZERO) - pay_before.get(number, ZERO)
        end = start + ch_in.get(number, ZERO) - pay_in.get(number, ZERO)
        vals = [start, ch_in.get(number, ZERO), pay_in.get(number, ZERO), end,
                adv.get(number, ZERO)]
        for i, v in enumerate(vals):
            tot[i] += v
        sh.values([number, owners.get(number, "")] + vals, money=(3, 4, 5, 6, 7))
    sh.values(["Итого", ""] + [tot[i] for i in range(5)], bold=True,
              money=(3, 4, 5, 6, 7))
    return _save(wb)
